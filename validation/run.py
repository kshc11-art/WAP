"""Local, read-only golden comparison. Never runs a program or promotes output to golden."""
from pathlib import Path
import argparse, csv, decimal, hashlib, json, sys
from datetime import datetime, timezone
from zipfile import BadZipFile
from xml.etree.ElementTree import ParseError
from contextlib import ExitStack

ROOT = Path(__file__).resolve().parent

class ValidationError(ValueError):
    pass

def strict_object(pairs):
    result = {}
    for key, value in pairs:
        if key in result:
            raise ValidationError(f'duplicate JSON property: {key}')
        result[key] = value
    return result

def read_json(path):
    return json.loads(Path(path).read_text(encoding='utf-8-sig'), object_pairs_hook=strict_object,
                      parse_float=decimal.Decimal)

def digest(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()

def within(folder, relative):
    if not isinstance(relative, str) or not relative:
        raise ValidationError('empty relative path')
    root = Path(folder).resolve()
    target = (root / relative).resolve()
    if target == root or not target.is_relative_to(root):
        raise ValidationError(f'path outside data folder: {relative}')
    return target

def atom(value):
    if value is None:
        return ['null']
    if isinstance(value, bool):
        return ['bool', value]
    if isinstance(value, (int, float, decimal.Decimal)):
        n = decimal.Decimal(str(value))
        if not n.is_finite():
            raise ValidationError('non-finite number')
        sign, digits, exponent = n.as_tuple()
        if not any(digits):
            return ['number', '0']
        digits = list(digits)
        while digits[-1] == 0:
            digits.pop(); exponent += 1
        return ['number', ('-' if sign else '') + ''.join(map(str, digits)) + 'e' + str(exponent)]
    if hasattr(value, 'isoformat'):
        return [type(value).__name__, value.isoformat()]
    if isinstance(value, dict):
        return ['object', [[k, atom(v)] for k, v in sorted(value.items())]]
    if isinstance(value, list):
        return ['array', [atom(v) for v in value]]
    return ['text', str(value)]

def canonical(value):
    return json.dumps(atom(value), ensure_ascii=False, separators=(',', ':'))

def table(headers, rows, keys=()):
    if not headers or any(h is None or str(h).strip() == '' for h in headers):
        raise ValidationError('empty table/header')
    headers = [str(h) for h in headers]
    if len(headers) != len(set(headers)):
        raise ValidationError('duplicate column names')
    if len(keys) != len(set(keys)) or any(k not in headers for k in keys):
        raise ValidationError('missing/duplicate key columns')
    result = {}
    for index, values in enumerate(rows, 1):
        if len(values) != len(headers):
            raise ValidationError(f'row width mismatch: row {index}')
        record = dict(zip(headers, values))
        if keys:
            vals = [record[k] for k in keys]
            if any(v is None or (isinstance(v, str) and not v.strip()) for v in vals):
                raise ValidationError(f'empty row key: row {index}')
            key = canonical(vals)
        else:
            key = str(index)
        if key in result:
            raise ValidationError(f'duplicate row key: row {index}')
        result[key] = {h: atom(v) for h, v in record.items()}
    return {'headers': headers, 'rows': result}

def load_table(path, spec):
    suffix = Path(path).suffix.lower()
    keys = spec.get('keys', [])
    if suffix in ('.csv', '.tsv'):
        try:
            with open(path, encoding=spec.get('encoding', 'utf-8-sig'), newline='') as stream:
                rows = list(csv.reader(stream, delimiter='\t' if suffix == '.tsv' else ',', strict=True))
        except csv.Error as e:
            raise ValidationError('malformed CSV/TSV: '+str(e)) from e
        if not rows:
            raise ValidationError('empty CSV/TSV')
        return {'table': table(rows[0], rows[1:], keys)}
    if suffix == '.json':
        value = read_json(path)
        if keys:
            if not isinstance(value, list) or not value or not all(isinstance(r, dict) for r in value):
                raise ValidationError('keyed JSON must be a nonempty array of objects')
            headers = sorted(value[0])
            if any(sorted(r) != headers for r in value):
                raise ValidationError('JSON row schema mismatch')
            return {'table': table(headers, [[r[h] for h in headers] for r in value], keys)}
        return {'json': atom(value)}
    if suffix == '.xlsx':
        try:
            import openpyxl
        except ImportError as e:
            raise ValidationError('XLSX comparison requires openpyxl; install requirements-dev.txt') from e
        with ExitStack() as resources:
            # Own the input handles even if openpyxl fails before returning a workbook.
            values = openpyxl.load_workbook(resources.enter_context(open(path, 'rb')), data_only=True, read_only=True)
            resources.callback(values.close)
            formulas = openpyxl.load_workbook(resources.enter_context(open(path, 'rb')), data_only=False, read_only=True)
            resources.callback(formulas.close)
            result = {'sheet_order': values.sheetnames, 'sheets': {}}
            selected = spec.get('sheets')
            if selected is not None and (not selected or set(selected) != set(values.sheetnames)):
                raise ValidationError('configured sheets must cover the complete workbook')
            for name in values.sheetnames:
                settings = selected[name] if selected else {}
                # Some producers leave stale worksheet dimensions after adding cells.
                # Parse the stored cells rather than silently truncating to that hint.
                values[name].reset_dimensions()
                formulas[name].reset_dimensions()
                rows = []
                for vr, fr in zip(values[name].iter_rows(), formulas[name].iter_rows()):
                    row = []
                    for v, f in zip(vr, fr):
                        if f.data_type == 'f' and v.value is None:
                            raise ValidationError(f'formula has no cached result: {name}!{f.coordinate}')
                        if f.data_type == 'e' or v.data_type == 'e':
                            raise ValidationError(f'Excel error: {name}!{f.coordinate}')
                        row.append(v.value)
                    rows.append(row)
                # Ignore formatting-only trailing cells, but not internal empty rows.
                while rows and all(v is None for v in rows[-1]):
                    rows.pop()
                width = max((i + 1 for r in rows for i, v in enumerate(r) if v is not None), default=0)
                rows = [r[:width] + [None] * max(0, width-len(r)) for r in rows]
                if settings.get('keys'):
                    h = settings.get('header_row', 1)
                    if not isinstance(h, int) or h < 1 or h > len(rows):
                        raise ValidationError('invalid header_row')
                    result['sheets'][name] = {'preamble': atom(rows[:h-1]), 'table': table(rows[h-1], rows[h:], settings['keys'])}
                else:
                    result['sheets'][name] = atom(rows)
            return result
    if keys:
        raise ValidationError('row keys unsupported for this file type')
    return {'binary_sha256': digest(path)}

def compare_files(actual, golden, spec):
    if Path(actual).suffix.lower() != Path(golden).suffix.lower():
        raise ValidationError('file type mismatch')
    a, g = load_table(actual, spec), load_table(golden, spec)
    changes = []
    def walk(x, y, location='$'):
        if type(x) != type(y):
            changes.append({'location': location, 'kind': 'type'}); return
        if isinstance(x, dict):
            for k in sorted(set(x) | set(y)):
                pos = location + '/' + k
                if k not in x:
                    changes.append({'location': pos, 'kind': 'missing-actual'})
                elif k not in y:
                    changes.append({'location': pos, 'kind': 'extra-actual'})
                else:
                    walk(x[k], y[k], pos)
        elif isinstance(x, list):
            if len(x) != len(y):
                changes.append({'location': location, 'kind': 'length', 'actual': len(x), 'golden': len(y)})
            for i, (left, right) in enumerate(zip(x, y)):
                walk(left, right, location + '/' + str(i))
        elif x != y:
            changes.append({'location': location, 'kind': 'value'})
    walk(a, g)
    return {'status': 'PASS' if not changes else 'FAIL', 'difference_count': len(changes), 'differences': changes[:200],
            'actual_sha256': digest(actual), 'golden_sha256': digest(golden)}

def compare_case(case, root=ROOT):
    name = case['id']
    if not re_id(name):
        raise ValidationError('invalid case id')
    if case.get('status') != 'ready' or not case.get('raw') or not case.get('comparisons'):
        return {'id': name, 'status': 'WAITING', 'reason': 'raw/golden comparison specification not ready'}
    rawroot, goldroot, actroot = [root/k/name for k in ('raw','golden','actual')]
    raw_seen = set()
    for item in case['raw']:
        if item['path'] in raw_seen:
            raise ValidationError('duplicate raw path')
        raw_seen.add(item['path'])
        p = within(rawroot, item['path'])
        if not p.is_file():
            return {'id': name, 'status': 'WAITING', 'reason': 'raw file missing: '+item['path']}
        if not item.get('sha256') or digest(p) != item['sha256']:
            raise ValidationError('raw hash absent or changed: '+item['path'])
    if {p.relative_to(rawroot).as_posix() for p in rawroot.rglob('*') if p.is_file()} != raw_seen:
        raise ValidationError('raw contains unregistered input files')
    outputs = []
    seen = set()
    for item in case['comparisons']:
        filename = item['path']
        if filename in seen:
            raise ValidationError('duplicate comparison path')
        seen.add(filename)
        a, g = within(actroot, filename), within(goldroot, filename)
        if not a.is_file() or not g.is_file():
            return {'id': name, 'status': 'WAITING', 'reason': 'actual/golden file missing: '+filename}
        if not item.get('golden_sha256') or digest(g) != item['golden_sha256']:
            raise ValidationError('golden hash absent or changed: '+filename)
        outputs.append({'file': filename, **compare_files(a, g, item)})
    actual_set = {p.relative_to(actroot).as_posix() for p in actroot.rglob('*') if p.is_file()}
    golden_set = {p.relative_to(goldroot).as_posix() for p in goldroot.rglob('*') if p.is_file()}
    if actual_set != seen or golden_set != seen:
        raise ValidationError('actual/golden contains unregistered or missing output files')
    return {'id': name, 'status': 'PASS' if all(x['status']=='PASS' for x in outputs) else 'FAIL', 'files': outputs}

def re_id(name):
    import re
    return isinstance(name, str) and re.fullmatch(r'[a-z0-9][a-z0-9-]*', name)

def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('command', choices=['status','compare'], nargs='?', default='status')
    parser.add_argument('case', nargs='?')
    args = parser.parse_args(argv)
    try:
        config = read_json(ROOT/'cases.json')
        cases = config['cases']
        ids = [x['id'] for x in cases]
        if not cases or len(ids) != len(set(ids)) or not all(re_id(x) for x in ids):
            raise ValidationError('empty/invalid case registry')
        if args.command == 'status':
            print('Actual business validation: pending until raw + golden + actual are supplied.')
            for case in cases:
                print(case['id'] + ': ' + case['status'])
            return 2  # Configuration status is never evidence of an executed comparison.
        if args.case not in ids:
            raise ValidationError('choose a registered case id')
        result = compare_case(next(x for x in cases if x['id']==args.case), root=ROOT)
        save_report(args.case, result)
        return {'PASS':0,'FAIL':1,'WAITING':2}[result['status']]
    except (ValidationError, OSError, ValueError, KeyError, TypeError, BadZipFile, ParseError) as e:
        if args.command == 'compare' and re_id(args.case):
            save_report(args.case, {'id': args.case, 'status': 'ERROR', 'reason': str(e)})
        print('ERROR: '+str(e), file=sys.stderr)
        return 1

def save_report(case_id, result):
    result['checked_at_utc'] = datetime.now(timezone.utc).isoformat()
    report = ROOT/'reports'/(case_id+'.json')
    report.parent.mkdir(exist_ok=True)
    report.write_text(json.dumps(result, ensure_ascii=False, indent=2)+'\n', encoding='utf-8')
    print(json.dumps(result, ensure_ascii=False, indent=2))

if __name__ == '__main__':
    sys.exit(main())
