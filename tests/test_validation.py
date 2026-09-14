import importlib.util, json, tempfile, unittest
from pathlib import Path
R=Path(__file__).resolve().parents[1]
s=importlib.util.spec_from_file_location('comparison', R/'validation/run.py')
m=importlib.util.module_from_spec(s);s.loader.exec_module(m)

class ComparisonTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory();self.root=Path(self.temp.name)
    def tearDown(self):self.temp.cleanup()
    def pair(self,a,g,suffix='.csv',spec=None):
        ap=self.root/('actual'+suffix);gp=self.root/('golden'+suffix)
        ap.write_text(a,encoding='utf-8');gp.write_text(g,encoding='utf-8')
        return m.compare_files(ap,gp,spec or {})
    def test_keyed_order(self):
        self.assertEqual(self.pair('id,n\nB,2\nA,1\n','id,n\nA,1\nB,2\n',spec={'keys':['id']})['status'],'PASS')
    def test_order_without_key(self):
        self.assertEqual(self.pair('id\nB\nA\n','id\nA\nB\n')['status'],'FAIL')
    def test_same_total_different_rows(self):
        self.assertEqual(self.pair('id,n\nA,2\nB,1\n','id,n\nA,1\nB,2\n',spec={'keys':['id']})['status'],'FAIL')
    def test_duplicate_keys(self):
        with self.assertRaises(m.ValidationError):self.pair('id\nA\nA\n','id\nA\n',spec={'keys':['id']})
    def test_blank_keys(self):
        with self.assertRaises(m.ValidationError):self.pair('id,n\n ,1\n','id,n\nA,1\n',spec={'keys':['id']})
    def test_duplicate_headers(self):
        with self.assertRaises(m.ValidationError):self.pair('id,id\nA,B\n','id\nA\n')
    def test_missing_column(self):
        with self.assertRaises(m.ValidationError):self.pair('n\n1\n','id\nA\n',spec={'keys':['id']})
    def test_leading_zero(self):
        self.assertEqual(self.pair('id\n01\n','id\n1\n')['status'],'FAIL')
    def test_json_null_zero(self):
        self.assertEqual(self.pair('{"x":null}','{"x":0}',suffix='.json')['status'],'FAIL')
    def test_json_bool_number(self):
        self.assertEqual(self.pair('{"x":true}','{"x":1}',suffix='.json')['status'],'FAIL')
    def test_json_duplicate_property(self):
        with self.assertRaises(m.ValidationError):self.pair('{"x":1,"x":2}','{"x":2}',suffix='.json')
    def test_json_precision(self):
        self.assertEqual(self.pair('0.1234567890123456789012345678901','0.1234567890123456789012345678902',suffix='.json')['status'],'FAIL')
        self.assertEqual(self.pair('1234567890123456789012345678901','1234567890123456789012345678902',suffix='.json')['status'],'FAIL')
        self.assertEqual(self.pair('1.00','1',suffix='.json')['status'],'PASS')
    def test_error_replaces_old_report(self):
        from unittest.mock import patch
        (self.root/'reports').mkdir();old=self.root/'reports/example.json'
        old.write_text('{"status":"PASS"}')
        (self.root/'cases.json').write_text('{"cases":[]}')
        with patch.object(m,'ROOT',self.root):
            self.assertEqual(m.main(['compare','example']),1)
        self.assertEqual(json.loads(old.read_text())['status'],'ERROR')
    def test_empty_csv(self):
        with self.assertRaises(m.ValidationError):self.pair('','')
    def test_path_escape(self):
        with self.assertRaises(m.ValidationError):m.within(self.root,'../outside')
    def test_pending(self):
        self.assertEqual(m.compare_case({'id':'example','status':'awaiting-data'},self.root)['status'],'WAITING')
    def test_hash_and_extra_file(self):
        for kind in ('raw','golden','actual'):(self.root/kind/'example').mkdir(parents=True)
        r=self.root/'raw/example/input.json';r.write_text('{}')
        a=self.root/'actual/example/report.csv';g=self.root/'golden/example/report.csv'
        a.write_text('id,n\nA,1\n');g.write_bytes(a.read_bytes())
        case={'id':'example','status':'ready','raw':[{'path':'input.json','sha256':m.digest(r)}],
              'comparisons':[{'path':'report.csv','golden_sha256':m.digest(g),'keys':['id']}]}
        self.assertEqual(m.compare_case(case,self.root)['status'],'PASS')
        r.write_text('{"changed":true}')
        with self.assertRaises(m.ValidationError):m.compare_case(case,self.root)
        case['raw'][0]['sha256']=m.digest(r)
        (a.parent/'extra.csv').write_text('id\nB\n')
        with self.assertRaises(m.ValidationError):m.compare_case(case,self.root)
        (a.parent/'extra.csv').unlink();g.write_text('id,n\nA,2\n')
        with self.assertRaises(m.ValidationError):m.compare_case(case,self.root)
    def test_xlsx_cells_and_cache(self):
        import openpyxl
        a=self.root/'a.xlsx';g=self.root/'g.xlsx'
        wb=openpyxl.Workbook();ws=wb.active;ws.append(['id','value']);ws.append(['A',1]);wb.save(a);wb.save(g)
        self.assertEqual(m.compare_files(a,g,{})['status'],'PASS')
        ws['B2']=2;wb.save(a)
        self.assertEqual(m.compare_files(a,g,{})['status'],'FAIL')
        ws['B2']='=1+1';wb.save(a)
        with self.assertRaises(m.ValidationError):m.compare_files(a,g,{})
    def test_xlsx_missing_sheet(self):
        import openpyxl
        a=self.root/'a.xlsx';g=self.root/'g.xlsx';wb=openpyxl.Workbook();wb.active.append(['value']);wb.save(a)
        wb.create_sheet('another');wb.save(g)
        self.assertEqual(m.compare_files(a,g,{})['status'],'FAIL')

if __name__=='__main__':unittest.main()
